#!/usr/bin/env python3
"""CUMCM 求解交付与优化纪律审计（Type-A 机器门禁）。

用法:
  python solution_audit.py <工作区根目录> [--mode deliverable|optimization|all]

读取两份机器可读清单（均由步骤 1/3 生成，按题填写，不内置任何赛题答案）:
  plan/solution_manifest.json   全链路矩阵（字段见模板 PROBLEM_ANALYSIS.md 第 8 节表格 + 第 11 节 JSON 字段清单）
  code/run_registry.json        运行台账（字段见 references/03-编程实现.md）

--mode deliverable（每问完整性 + 官方附件）:
  1. manifest 存在且合法，subproblems 非空
  2. 每问: 全模型入口存在；结果文件存在且非空
  3. 若声明 official_template: 结果文件 != 原始模板；xlsx 则再比对 sheet 名与
     首行表头/列序（大小写与首尾空格归一化后逐列相等）
  4. 若 registry 存在: 每问最新一次运行须 status=ok 且 result.feasible=True
     （manifest 中该问声明了 hard_constraints 时生效）

--mode optimization（优化纪律）:
  1. rated_targets 非空（优化题必须有额定/硬指标）
  2. 声明 surrogate 的问题必须给出 surrogate_terms（物理项清单）且
     calibration_samples > 0
  3. 声明 hard_constraints 的问题在 registry 中必须存在一次
     status=ok 且 result.feasible=True 的运行（存在最佳可行解）
  4. 范式切换: 若某问存在 >= N 轮（默认 3，可配）连续不可行运行，且此后/最后的
     运行仍未切换 paradigm -> FAIL
  5. 预算: 全部运行耗时总和 > budget_total_min -> FAIL；单问超 budget_min -> WARN
  6. aborted/empty/failed 运行存在 -> WARN（即使后续修复，也提示人工复核）

退出码: 0 = PASS（可含 WARN）；1 = 存在 FAIL。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

# 与 PROBLEM_ANALYSIS.md / SOLUTION_ACCEPTANCE_CONTRACT.md 保持一致的可配默认值
DEFAULT_THRESHOLDS = {
    "surrogate_rank_corr": 0.9,
    "surrogate_delta_pct": 0.02,
    "infeasible_rounds_before_switch": 3,
    "later_question_budget_pct": 0.2,
}
FAIL_STATUSES = {"aborted", "empty", "failed"}


def load_json(path: Path):
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as e:  # noqa: BLE001
        return {"__parse_error__": str(e)}


def rel(p: Path, ws: Path) -> str:
    try:
        return str(p.relative_to(ws))
    except ValueError:
        return str(p)


def check_file(ws: Path, rel_path, label: str, results: list):
    p = (ws / rel_path) if rel_path else None
    if not rel_path:
        results.append(("FAIL", f"{label}: manifest 未声明路径"))
        return None
    if not p.exists():
        results.append(("FAIL", f"{label}: 文件不存在 {rel(p, ws)}"))
        return None
    if p.stat().st_size == 0:
        results.append(("FAIL", f"{label}: 文件为空 {rel(p, ws)}"))
        return None
    return p


def xlsx_headers(path: Path):
    """返回 (sheet名列表, 首个sheet首行表头列表)；非 xlsx 或读取失败返回 None。"""
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return None
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = wb.sheetnames
        ws = wb[sheets[0]]
        header = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            header = [str(c).strip().lower() if c is not None else "" for c in row]
        wb.close()
        return sheets, header
    except Exception:  # noqa: BLE001
        return None


def norm(s: str) -> str:
    return "".join(str(s).split()).lower()


def audit_deliverable(ws: Path, manifest: dict) -> list:
    results = []
    subs = manifest.get("subproblems") or []
    if not subs:
        results.append(("FAIL", "manifest.subproblems 为空（全链路矩阵缺失）"))
        return results

    registry = load_json(ws / "code" / "run_registry.json")
    registry_by_sub = {}
    if isinstance(registry, list):
        for r in registry:
            registry_by_sub.setdefault(r.get("subproblem", r.get("run_id", "?")), []).append(r)

    for sp in subs:
        sid = sp.get("id", "?")
        n_before = len(results)
        full = check_file(ws, sp.get("full_model"), f"[{sid}] 全模型入口", results)
        res = check_file(ws, sp.get("result_file"), f"[{sid}] 结果文件", results)

        tpl_path = sp.get("official_template")
        if tpl_path:
            tpl = (ws / tpl_path)
            if not tpl.exists():
                results.append(("FAIL", f"[{sid}] 官方模板不存在 {tpl_path}"))
            elif res is not None:
                if res.read_bytes() == tpl.read_bytes():
                    results.append(("FAIL", f"[{sid}] 结果文件与原始模板完全相同（疑似未填）"))
                else:
                    th = xlsx_headers(tpl)
                    rh = xlsx_headers(res)
                    if th and rh:
                        if th[0] != rh[0]:
                            results.append(("FAIL", f"[{sid}] sheet 名不一致：模板 {th[0]} vs 结果 {rh[0]}"))
                        elif th[1] != rh[1]:
                            results.append(("FAIL", f"[{sid}] 首行表头/列序不一致：\n    模板 {th[1]}\n    结果 {rh[1]}"))
        # 硬约束：由 registry 最新运行判定（manifest 声明 hard_constraints 时生效）
        runs = registry_by_sub.get(sid, [])
        if sp.get("hard_constraints") and runs:
            last = runs[-1]
            ok = last.get("status") == "ok"
            feasible = (last.get("result") or {}).get("feasible")
            if not (ok and feasible is True):
                results.append(("FAIL", f"[{sid}] 最新运行未满足硬约束（status={last.get('status')}, feasible={feasible}）"))
        if len(results) == n_before:
            results.append(("PASS", f"[{sid}] 全模型入口、结果文件、官方附件与硬约束均通过"))
    return results


def audit_optimization(ws: Path, manifest: dict) -> list:
    results = []
    rated = manifest.get("rated_targets") or {}
    if not rated:
        results.append(("FAIL", "优化题必须有 rated_targets（额定/硬指标），见 manifest"))
    th = {**DEFAULT_THRESHOLDS, **(manifest.get("thresholds") or {})}
    n_switch = int(th.get("infeasible_rounds_before_switch", 3))

    registry = load_json(ws / "code" / "run_registry.json")
    if not isinstance(registry, list) or not registry:
        results.append(("FAIL", "code/run_registry.json 缺失或为空（优化纪律无法审计）"))
        return results

    subs = manifest.get("subproblems") or []
    if not subs:
        results.append(("FAIL", "manifest.subproblems 为空"))
        return results

    by_sub = {}
    for r in registry:
        by_sub.setdefault(r.get("subproblem", "?"), []).append(r)

    for sp in subs:
        sid = sp.get("id", "?")
        runs = by_sub.get(sid, [])
        # 代理保真度
        if sp.get("surrogate"):
            terms = sp.get("surrogate_terms") or []
            ok_terms = bool(terms)
            ok_cal = bool(sp.get("calibration_samples"))
            ok_file = (ws / sp["surrogate"]).exists()
            if not ok_terms:
                results.append(("FAIL", f"[{sid}] 声明了代理但 surrogate_terms（物理项清单）为空"))
            if not ok_cal:
                results.append(("FAIL", f"[{sid}] 声明了代理但 calibration_samples 缺失或为 0"))
            if not ok_file:
                results.append(("FAIL", f"[{sid}] 代理入口文件不存在 {sp['surrogate']}"))
            if ok_terms and ok_cal and ok_file:
                results.append(("PASS", f"[{sid}] 代理物理项清单({len(terms)}项)与校准样本({sp.get('calibration_samples')})齐全"))
        # 存在最佳可行解
        if sp.get("hard_constraints"):
            feasible_runs = [r for r in runs
                             if r.get("status") == "ok"
                             and (r.get("result") or {}).get("feasible") is True]
            if not feasible_runs:
                results.append(("FAIL", f"[{sid}] 声明了硬约束但无任何可行运行（最佳可行解缺失）"))
            else:
                results.append(("PASS", f"[{sid}] 存在 {len(feasible_runs)} 条可行运行（最佳可行解存在）"))
        # 范式切换：连续 N 轮不可行后必须切换 paradigm
        streak = 0
        max_streak = 0
        streak_paradigm = None
        last_paradigm = None
        for r in runs:
            st = r.get("status")
            feasible = (r.get("result") or {}).get("feasible")
            para = r.get("paradigm", "")
            if para:
                last_paradigm = para
            if st == "ok" and feasible is False:
                if streak == 0:
                    streak_paradigm = para
                streak += 1
                max_streak = max(max_streak, streak)
            else:
                streak = 0
        if max_streak >= n_switch and last_paradigm == streak_paradigm and streak_paradigm:
            results.append(("FAIL", f"[{sid}] 连续 {max_streak} 轮不可行（>= {n_switch}）但未切换范式（仍为 '{streak_paradigm}'）"))
        elif max_streak >= n_switch:
            results.append(("PASS", f"[{sid}] 连续 {max_streak} 轮不可行后已切换范式（{streak_paradigm} -> {last_paradigm}）"))
        else:
            results.append(("PASS", f"[{sid}] 无连续 {n_switch} 轮不可行（最大连续 {max_streak}）"))
        # 空/中断/失败运行提示
        bad = [r for r in runs if r.get("status") in FAIL_STATUSES]
        if bad:
            results.append(("WARN", f"[{sid}] 存在 {len(bad)} 条 {sorted(set(r['status'] for r in bad))} 运行，请人工确认未被误判为完成"))

    # 预算
    total_dur = sum(int(r.get("duration_s") or 0) for r in registry)
    budget_total = manifest.get("budget_total_min")
    if budget_total:
        if total_dur > budget_total * 60:
            results.append(("FAIL", f"总运行耗时 {total_dur / 60:.1f} min 超过预算 {budget_total} min"))
        else:
            results.append(("PASS", f"总运行耗时 {total_dur / 60:.1f} min 在预算 {budget_total} min 内"))
    for sp in subs:
        sid = sp.get("id", "?")
        bmin = sp.get("budget_min")
        if bmin:
            used = sum(int(r.get("duration_s") or 0) for r in by_sub.get(sid, []))
            if used > bmin * 60:
                results.append(("WARN", f"[{sid}] 单问耗时 {used / 60:.1f} min 超过该问预算 {bmin} min"))
    return results


def main() -> int:
    ap = argparse.ArgumentParser(description="CUMCM solution delivery & optimization audit.")
    ap.add_argument("workspace", help="工作区根目录")
    ap.add_argument("--mode", choices=["deliverable", "optimization", "all"], default="all")
    a = ap.parse_args()

    ws = Path(a.workspace)
    if not ws.is_dir():
        print(f"FAIL: 工作区不存在 {ws}")
        return 1

    manifest = load_json(ws / "plan" / "solution_manifest.json")
    if manifest is None:
        print("FAIL: plan/solution_manifest.json 缺失（步骤 1 必须生成全链路矩阵）")
        return 1
    if "__parse_error__" in manifest:
        print(f"FAIL: manifest JSON 解析失败: {manifest['__parse_error__']}")
        return 1

    results = []
    if a.mode in ("deliverable", "all"):
        results += audit_deliverable(ws, manifest)
    if a.mode in ("optimization", "all"):
        results += audit_optimization(ws, manifest)

    n_fail = 0
    n_warn = 0
    for level, msg in results:
        print(f"{level}: {msg}")
        if level == "FAIL":
            n_fail += 1
        elif level == "WARN":
            n_warn += 1
    print("-" * 60)
    print(f"RESULT: {'PASS' if n_fail == 0 else 'FAIL'}  ({len(results)} 项, {n_fail} FAIL, {n_warn} WARN)")
    return 0 if n_fail == 0 else 1


if __name__ == "__main__":
    sys.exit(main())