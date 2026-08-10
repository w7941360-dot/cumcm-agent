#!/usr/bin/env python3
"""CUMCM 输入处理：赛题 PDF 提取文本 + 原始数据表格(xlsx)探查。

用法:
  python process_inputs.py <工作区根目录> <赛题.pdf> [数据1.xlsx 数据2.xlsx ...]

产出:
  <工作区>/problem/<赛题名>_提取文本.txt    PDF 全文（分页）
  <工作区>/problem/DATA_PROFILE.md          每个 xlsx 的数据探查报告（sheet/行列/列名/缺失/类型/统计）
  <工作区>/data/                             原始数据文件副本（后续清洗后的数据也放这里）

依赖（自动探测，任一可用即可）:
  PDF:  pymupdf(fitz) / pypdf / PyPDF2
  xlsx: openpyxl（两个 Python 环境都已验证可用）
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path
from datetime import datetime


# ---------- PDF 文本提取 ----------

def extract_pdf_text(pdf_path: Path):
    """尝试 fitz -> pypdf -> PyPDF2，返回 (全文, 用的库)；全失败返回 (None, None)。"""
    parts = []

    def wrap_page(i, text):
        return f"===== 第 {i + 1} 页 =====\n{text or ''}"

    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(pdf_path))
        for i, page in enumerate(doc):
            parts.append(wrap_page(i, page.get_text()))
        doc.close()
        return "\n".join(parts), "PyMuPDF(fitz)"
    except Exception:
        pass

    try:
        import pypdf
        reader = pypdf.PdfReader(str(pdf_path))
        for i, page in enumerate(reader.pages):
            parts.append(wrap_page(i, page.extract_text() or ""))
        return "\n".join(parts), "pypdf"
    except Exception:
        pass

    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(str(pdf_path))
        for i, page in enumerate(reader.pages):
            parts.append(wrap_page(i, page.extract_text() or ""))
        return "\n".join(parts), "PyPDF2"
    except Exception:
        return None, None


# ---------- xlsx 探查 ----------

def _guess_type(v):
    if v is None:
        return "缺失"
    if isinstance(v, bool):
        return "布尔"
    if isinstance(v, (int, float)):
        return "数值"
    if isinstance(v, datetime):
        return "日期"
    return "文本"


def profile_xlsx(xlsx_path: Path) -> str:
    """用 openpyxl 探查每个 sheet，输出 markdown 报告。"""
    lines = [f"# 数据探查：{xlsx_path.name}", ""]
    try:
        from openpyxl import load_workbook
    except Exception as e:
        lines.append(f"⚠️ 无法读取（缺少 openpyxl：{e}）")
        return "\n".join(lines)

    wb = None
    try:
        wb = load_workbook(str(xlsx_path), data_only=True)
    except Exception as e:
        lines.append(f"⚠️ 打开失败：{e}")
        return "\n".join(lines)

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            lines.append(f"\n## Sheet: {ws.title}（空表）")
            continue
        header = rows[0]
        n_col = len(header)
        n_row = len(rows) - 1
        lines.append(f"\n## Sheet: {ws.title}")
        lines.append(f"- 维度：{n_row} 行数据 x {n_col} 列")
        lines.append(f"- 列名：{', '.join(str(c) if c is not None else '(无)' for c in header)}")
        lines.append("- 前 3 行样例：")
        for r in rows[1:4]:
            lines.append("  - " + " | ".join(str(c) if c is not None else "" for c in r))
        # 逐列统计
        lines.append("- 逐列统计：")
        for j in range(n_col):
            col_vals = [r[j] for r in rows[1:] if j < len(r)]
            n_missing = sum(1 for v in col_vals if v is None or str(v).strip() == "")
            n_valid = len(col_vals) - n_missing
            types = {}
            for v in col_vals:
                if v is None or str(v).strip() == "":
                    continue
                t = _guess_type(v)
                types[t] = types.get(t, 0) + 1
            type_desc = ", ".join(f"{k}={v}" for k, v in sorted(types.items()))
            # 数值列统计
            num_vals = [v for v in col_vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
            stat = ""
            if num_vals:
                stat = f" | min={min(num_vals):g} max={max(num_vals):g} mean={sum(num_vals)/len(num_vals):.4g}"
            lines.append(
                f"  - 列{j+1}「{header[j] if j < len(header) and header[j] is not None else ''}」"
                f"：有效 {n_valid}/{len(col_vals)}，缺失 {n_missing}，类型({type_desc}){stat}"
            )
    if wb is not None:
        wb.close()
    return "\n".join(lines)


# ---------- 主流程 ----------

def main() -> int:
    ap = argparse.ArgumentParser(description="CUMCM input processor: PDF text + xlsx profile.")
    ap.add_argument("workspace", help="工作区根目录")
    ap.add_argument("pdf", help="赛题 PDF 路径")
    ap.add_argument("xlsx", nargs="*", default=[], help="原始数据 xlsx 路径（可多个）")
    a = ap.parse_args()

    ws = Path(a.workspace)
    problem_dir = ws / "problem"
    data_dir = ws / "data"
    problem_dir.mkdir(parents=True, exist_ok=True)
    data_dir.mkdir(parents=True, exist_ok=True)

    # 1) 赛题 PDF
    pdf = Path(a.pdf)
    if not pdf.is_file():
        print(f"ERROR: 赛题 PDF 不存在: {pdf}")
        return 1
    shutil.copy2(pdf, problem_dir / pdf.name)
    text, lib = extract_pdf_text(pdf)
    if text is None:
        print("WARN: PDF 无法提取文本（可能是扫描版/图片版，无文字层）。")
        print("      → 处理办法：a) 用 OCR 识别；b) 把 PDF 里每页截图，人工/看图工具读取；c) 用视觉模型读图。")
        print("      已把 PDF 原文件复制到 problem/，请按上述方式补充文本后继续。")
    else:
        out_txt = problem_dir / f"{pdf.stem}_提取文本.txt"
        out_txt.write_text(text, encoding="utf-8")
        print(f"OK: PDF 文本已提取（{lib}）→ {out_txt}（{len(text)} 字符）")
        n_pages = text.count("===== 第 ")
        empty_pages = [i + 1 for i, seg in enumerate(text.split("===== 第 ")[1:]) if not seg.split("=====", 1)[-1].strip()]
        if empty_pages:
            print(f"NOTE: 以下页无文字（可能含图/表/扫描）：{empty_pages}，请人工查看这些页")

    # 2) 数据 xlsx
    for x in a.xlsx:
        xp = Path(x)
        if not xp.is_file():
            print(f"WARN: 数据文件不存在，跳过: {xp}")
            continue
        shutil.copy2(xp, data_dir / xp.name)
        rep = profile_xlsx(xp)
        profile_path = problem_dir / "DATA_PROFILE.md"
        if profile_path.exists():
            profile_path.write_text(profile_path.read_text(encoding="utf-8") + "\n\n---\n\n" + rep, encoding="utf-8")
        else:
            profile_path.write_text(rep, encoding="utf-8")
        print(f"OK: 数据探查完成 → {profile_path}（{xp.name}）")

    print("提示：阅读 problem/赛题提取文本.txt 与 problem/DATA_PROFILE.md 后再进行赛题分析（步骤1）。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
